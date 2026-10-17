"""文件类型识别、结构验证及随机键附件存储。"""

from __future__ import annotations

import hashlib
import os
import re
import tempfile
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Literal
from urllib.parse import quote

import anyio
from fastapi import UploadFile
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader

ParseIntent = Literal["contract_pdf", "receipt_image", "order_xlsx"]

MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_PDF_BYTES = 20 * 1024 * 1024
MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 20
MAX_XLSX_ROWS = 500
MAX_IMAGE_PIXELS = 40_000_000
MAX_IMAGE_SIDE = 16_000
MAX_ZIP_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 100
CHUNK_BYTES = 64 * 1024

_ALLOWED = {
    "contract_pdf": ({".pdf"}, MAX_PDF_BYTES),
    "receipt_image": ({".jpg", ".jpeg", ".png"}, MAX_IMAGE_BYTES),
    "order_xlsx": ({".xlsx"}, MAX_XLSX_BYTES),
}
_MEDIA_TYPES = {
    "pdf": "application/pdf",
    "jpeg": "image/jpeg",
    "png": "image/png",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_STORAGE_KEY = re.compile(r"^[a-f0-9]{32}$")


class FileUploadError(Exception):
    """可返回给调用方的安全文件校验失败，不包含临时路径。"""

    def __init__(self, code: str, message: str, status_code: int | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = (
            status_code
            if status_code is not None
            else (413 if code == "FILE_TOO_LARGE" else 400)
        )


@dataclass(frozen=True, slots=True)
class StoredUpload:
    original_name: str
    storage_key: str
    path: Path
    media_type: str
    byte_size: int
    sha256: str
    metadata: dict[str, object]


def safe_filename(filename: str | None) -> str:
    """只保存单个显示名称，移除路径、控制字符和响应头注入字符。"""
    value = (filename or "").replace("\\", "/")
    name = PurePosixPath(value).name
    name = "".join(
        char for char in name if char.isprintable() and char not in "\r\n\x00"
    )
    name = name.strip(" .")[:255]
    if not name or name in {".", ".."}:
        raise FileUploadError("FILE_UNSUPPORTED", "文件名无效")
    return name


def max_size_for_intent(intent: ParseIntent) -> int:
    return _ALLOWED[intent][1]


def content_disposition(filename: str) -> str:
    """按 RFC 5987 编码显示名；该值从不参与存储路径生成。"""
    fallback = re.sub(r"[^A-Za-z0-9._-]", "_", filename) or "attachment"
    return f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{quote(filename, safe='')}"


def storage_path(root: Path, storage_key: str) -> Path:
    """拒绝非随机十六进制键，且仅在私有对象目录内解析。"""
    if not _STORAGE_KEY.fullmatch(storage_key):
        raise FileUploadError("FILE_UNREADABLE", "附件存储记录无效", 503)
    object_root = (root / "objects").resolve()
    path = object_root / storage_key
    if path.parent != object_root:
        raise FileUploadError("FILE_UNREADABLE", "附件存储记录无效", 503)
    return path


def _validate_pdf(path: Path) -> dict[str, object]:
    with path.open("rb") as stream:
        if b"%PDF-" not in stream.read(1024):
            raise FileUploadError("FILE_UNSUPPORTED", "文件内容与 PDF 格式不匹配")
    try:
        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted:
            raise FileUploadError("FILE_UNREADABLE", "不支持加密 PDF")
        pages = len(reader.pages)
    except FileUploadError:
        raise
    except Exception as exc:
        # pypdf uses parser-specific exceptions that vary across supported versions.
        raise FileUploadError("FILE_UNREADABLE", "PDF 文件损坏或无法读取") from exc
    if pages < 1:
        raise FileUploadError("FILE_UNREADABLE", "PDF 不包含可读取页面")
    if pages > MAX_PDF_PAGES:
        raise FileUploadError("FILE_TOO_LARGE", f"PDF 页数不能超过 {MAX_PDF_PAGES} 页")
    return {"page_count": pages}


def _validate_image(path: Path, extension: str) -> dict[str, object]:
    expected = "PNG" if extension == ".png" else "JPEG"
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as image:
                if image.format != expected:
                    raise FileUploadError(
                        "FILE_UNSUPPORTED", "文件扩展名与图片内容不匹配"
                    )
                width, height = image.size
                if (
                    width < 1
                    or height < 1
                    or width > MAX_IMAGE_SIDE
                    or height > MAX_IMAGE_SIDE
                    or width * height > MAX_IMAGE_PIXELS
                ):
                    raise FileUploadError("FILE_TOO_LARGE", "图片尺寸超出安全处理范围")
                image.verify()
            with Image.open(path) as image:
                image.load()
    except FileUploadError:
        raise
    except (
        UnidentifiedImageError,
        OSError,
        ValueError,
        Image.DecompressionBombError,
        Image.DecompressionBombWarning,
    ) as exc:
        raise FileUploadError("FILE_UNREADABLE", "图片文件损坏或无法读取") from exc
    return {"width": width, "height": height}


def _validate_xlsx(path: Path) -> dict[str, object]:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 4096:
                raise FileUploadError("FILE_UNSUPPORTED", "XLSX 内部文件数量超出限制")
            names = [item.filename for item in infos]
            if len(set(names)) != len(names):
                raise FileUploadError("FILE_UNREADABLE", "XLSX 内部文件重复")
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise FileUploadError("FILE_UNSUPPORTED", "仅支持标准 XLSX 工作簿")
            total_size = 0
            for item in infos:
                total_size += item.file_size
                if item.file_size > MAX_ZIP_UNCOMPRESSED_BYTES:
                    raise FileUploadError(
                        "FILE_TOO_LARGE", "XLSX 单个内部文件解压后过大"
                    )
                if (
                    item.file_size / max(item.compress_size, 1)
                    > MAX_ZIP_COMPRESSION_RATIO
                ):
                    raise FileUploadError(
                        "FILE_UNSUPPORTED", "XLSX 压缩比例超出安全范围"
                    )
                if "vbaproject.bin" in item.filename.casefold():
                    raise FileUploadError(
                        "FILE_UNSUPPORTED", "不支持包含宏的 Excel 文件"
                    )
            if total_size > MAX_ZIP_UNCOMPRESSED_BYTES:
                raise FileUploadError("FILE_TOO_LARGE", "XLSX 解压后总大小超出安全范围")
            with archive.open("[Content_Types].xml") as content_type_file:
                content_types = content_type_file.read(1024 * 1024 + 1)
            if len(content_types) > 1024 * 1024:
                raise FileUploadError("FILE_TOO_LARGE", "XLSX 类型声明文件超出安全限制")
            if b"macroEnabled" in content_types:
                raise FileUploadError("FILE_UNSUPPORTED", "不支持包含宏的 Excel 文件")
    except FileUploadError:
        raise
    except (zipfile.BadZipFile, OSError, KeyError, RuntimeError) as exc:
        raise FileUploadError("FILE_UNREADABLE", "XLSX 文件损坏或无法读取") from exc

    workbook = None
    try:
        # 从已校验文件流读取，避免解析库根据随机 .part 临时后缀拒绝文件。
        with path.open("rb") as stream:
            workbook = load_workbook(
                stream, read_only=True, data_only=False, keep_links=False
            )
            row_count = 0
            sheet_count = 0
            for worksheet in workbook.worksheets:
                sheet_count += 1
                has_header = False
                for row in worksheet.iter_rows(values_only=True):
                    if not any(value is not None for value in row):
                        continue
                    if has_header:
                        row_count += 1
                        if row_count > MAX_XLSX_ROWS:
                            raise FileUploadError(
                                "FILE_TOO_LARGE",
                                f"XLSX 明细行数不能超过 {MAX_XLSX_ROWS} 行",
                            )
                    else:
                        has_header = True
    except FileUploadError:
        raise
    except Exception as exc:
        raise FileUploadError("FILE_UNREADABLE", "XLSX 工作簿损坏或无法读取") from exc
    finally:
        if workbook is not None:
            workbook.close()
    return {"sheet_count": sheet_count, "row_count": row_count}


def _validate_file(
    path: Path, filename: str, intent: ParseIntent
) -> tuple[str, dict[str, object]]:
    extension = Path(filename).suffix.casefold()
    allowed_extensions, _ = _ALLOWED[intent]
    if extension not in allowed_extensions:
        raise FileUploadError("FILE_UNSUPPORTED", "文件扩展名与处理类型不匹配")
    if path.stat().st_size == 0:
        raise FileUploadError("FILE_UNREADABLE", "不支持空文件")
    if intent == "contract_pdf":
        return _MEDIA_TYPES["pdf"], _validate_pdf(path)
    if intent == "receipt_image":
        key = "png" if extension == ".png" else "jpeg"
        return _MEDIA_TYPES[key], _validate_image(path, extension)
    with path.open("rb") as stream:
        if stream.read(4) != b"PK\x03\x04":
            raise FileUploadError("FILE_UNSUPPORTED", "文件内容与 XLSX 格式不匹配")
    return _MEDIA_TYPES["xlsx"], _validate_xlsx(path)


async def store_upload(
    upload: UploadFile,
    *,
    root: Path,
    intent: ParseIntent,
) -> StoredUpload:
    """按块写私有暂存文件、同步计算摘要并在格式验证后原子转入对象目录。"""
    filename = safe_filename(upload.filename)
    max_bytes = max_size_for_intent(intent)
    root.mkdir(parents=True, exist_ok=True)
    staging = root / ".staging"
    objects = root / "objects"
    staging.mkdir(mode=0o700, parents=True, exist_ok=True)
    objects.mkdir(mode=0o700, parents=True, exist_ok=True)
    os.chmod(staging, 0o700)
    os.chmod(objects, 0o700)

    fd, temp_name = tempfile.mkstemp(prefix="upload-", suffix=".part", dir=staging)
    os.chmod(temp_name, 0o600)
    temp_path = Path(temp_name)
    digest = hashlib.sha256()
    byte_size = 0
    storage_key = os.urandom(16).hex()
    final_path = storage_path(root, storage_key)
    try:
        os.close(fd)
        async with await anyio.open_file(temp_path, "wb") as destination:
            while chunk := await upload.read(CHUNK_BYTES):
                byte_size += len(chunk)
                if byte_size > max_bytes:
                    raise FileUploadError(
                        "FILE_TOO_LARGE", "文件大小超过该类型的上传限制"
                    )
                digest.update(chunk)
                await destination.write(chunk)
        media_type, details = await anyio.to_thread.run_sync(
            _validate_file, temp_path, filename, intent
        )
        os.replace(temp_path, final_path)
        return StoredUpload(
            original_name=filename,
            storage_key=storage_key,
            path=final_path,
            media_type=media_type,
            byte_size=byte_size,
            sha256=digest.hexdigest(),
            metadata={"parse_intent": intent, **details},
        )
    except BaseException:
        try:
            temp_path.unlink(missing_ok=True)
        except OSError:
            pass
        try:
            final_path.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def limits_for_intent(intent: ParseIntent) -> dict[str, int | None]:
    return {
        "max_bytes": max_size_for_intent(intent),
        "max_pages": MAX_PDF_PAGES if intent == "contract_pdf" else None,
        "max_rows": MAX_XLSX_ROWS if intent == "order_xlsx" else None,
    }
