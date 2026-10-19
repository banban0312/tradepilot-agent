"""读取安全校验后的 XLSX，按订单分组生成带单元格来源的候选草稿。"""

from __future__ import annotations

import math
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.drafts.types import (
    CandidateOrder,
    CandidateOrderItem,
    FieldEvidenceInput,
    FieldSource,
    candidate_amounts,
)
from app.models import Contract, ContractBalance, Customer, Project

MAX_IMPORT_ROWS = 500
TEMPLATE_HEADERS = (
    "订单分组",
    "订单类型",
    "合同编号",
    "项目",
    "地址",
    "联系人",
    "交期",
    "物料名称",
    "规格",
    "数量",
    "单位",
    "含税单价",
)

_HEADER_ALIASES = {
    "订单分组": "group_id",
    "分组编号": "group_id",
    "订单类型": "order_type",
    "合同编号": "contract_number",
    "合同号": "contract_number",
    "项目": "project",
    "项目名称": "project",
    "地址": "address",
    "收货地址": "address",
    "联系人": "contact",
    "交期": "delivery_date",
    "交货日期": "delivery_date",
    "物料名称": "raw_name",
    "物料": "raw_name",
    "规格": "specification",
    "数量": "quantity",
    "单位": "unit",
    "含税单价": "unit_price",
    "单价": "unit_price",
}
_REQUIRED_COLUMNS = set(_HEADER_ALIASES.values())
_GROUP_FIELDS = (
    "order_type",
    "contract_number",
    "project",
    "address",
    "contact",
    "delivery_date",
)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _cell_issue(
    code: str, field: str, message: str, location: dict[str, object]
) -> dict[str, object]:
    return {"code": code, "field": field, "message": message, "locator": location}


def _order_type(value: Any) -> str | None:
    normalized = _text(value)
    if normalized is None:
        return None
    return {
        "销售": "SALES",
        "采购": "PURCHASE",
        "SALES": "SALES",
        "PURCHASE": "PURCHASE",
    }.get(normalized.upper() if normalized.isascii() else normalized)


def _decimal_text(value: Any) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError("必须填写有效数字")  # noqa: TRY004 - 统一 Excel 单元格错误格式
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("必须填写有限数字")
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("必须填写有效数字") from exc
    if not number.is_finite() or number < 0 or number >= Decimal(1000000000000):
        raise ValueError("数字超出允许范围")
    if number != 0 and number.normalize().as_tuple().exponent < -6:
        raise ValueError("最多支持 6 位小数")
    return format(number, "f")


def _date_text(value: Any) -> date | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError as exc:
            raise ValueError("日期请使用 YYYY-MM-DD 格式") from exc
    raise ValueError("日期请使用 YYYY-MM-DD 格式")


async def _find_contract(
    session: AsyncSession, department_ids: frozenset[UUID], number: str
) -> tuple[Contract, Customer, Project] | None:
    rows = (
        await session.execute(
            select(Contract, Customer, Project)
            .join(Customer, Customer.id == Contract.customer_id)
            .join(Project, Project.id == Contract.project_id)
            .where(
                Contract.number == number,
                Contract.department_id.in_(department_ids),
            )
        )
    ).all()
    return rows[0] if len(rows) == 1 else None


async def preview_workbook(
    path: Path,
    *,
    artifact_id: UUID,
    session: AsyncSession,
    department_ids: frozenset[UUID],
) -> dict[str, object]:
    """产出只读预览；候选状态不会被视为可批准或可提交订单。"""
    source = path.open("rb")
    try:
        workbook = load_workbook(
            source, read_only=True, data_only=False, keep_links=False
        )
    except BaseException:
        source.close()
        raise
    groups: dict[tuple[str, str], dict[str, Any]] = {}
    issues: list[dict[str, object]] = []
    row_count = 0
    contract_cache: dict[str, tuple[Contract, Customer, Project] | None] = {}
    balance_cache: dict[UUID, ContractBalance | None] = {}
    imported_by_contract: dict[UUID, Decimal] = {}
    try:
        for worksheet in workbook.worksheets:
            if worksheet.title == "填报说明":
                continue
            rows = worksheet.iter_rows()
            headers: dict[str, int] | None = None
            header_row_number = 1
            for row in rows:
                values = [cell.value for cell in row]
                if not any(value is not None for value in values):
                    continue
                if headers is None:
                    header_row_number = row[0].row
                    headers = {}
                    duplicate_columns: set[str] = set()
                    for column, value in enumerate(values, start=1):
                        label = _text(value)
                        field = _HEADER_ALIASES.get(label or "")
                        if field is None:
                            continue
                        if field in headers:
                            duplicate_columns.add(field)
                        else:
                            headers[field] = column
                    missing = sorted(_REQUIRED_COLUMNS - set(headers))
                    for field in sorted(duplicate_columns):
                        issues.append(
                            _cell_issue(
                                "DUPLICATE_COLUMN",
                                field,
                                "模板中存在重复字段列，请保留一列。",
                                {
                                    "sheet": worksheet.title,
                                    "row": header_row_number,
                                    "column": headers[field],
                                    "cell": f"{get_column_letter(headers[field])}{header_row_number}",
                                },
                            )
                        )
                    for field in missing:
                        issues.append(
                            _cell_issue(
                                "MISSING_COLUMN",
                                field,
                                "模板缺少必需列，请下载最新模板后重新填写。",
                                {
                                    "sheet": worksheet.title,
                                    "row": header_row_number,
                                    "column": None,
                                    "cell": None,
                                },
                            )
                        )
                    if missing or duplicate_columns:
                        break
                    continue

                row_count += 1
                if row_count > MAX_IMPORT_ROWS:
                    raise ValueError("Excel 明细行数超过 500 行，请拆分后导入。")
                raw_group = (
                    values[headers["group_id"] - 1]
                    if headers["group_id"] <= len(values)
                    else None
                )
                group_id = _text(raw_group)
                if group_id is None:
                    group_id = f"（缺少分组）第 {row[0].row} 行"
                key = (worksheet.title, group_id)
                group = groups.setdefault(
                    key,
                    {
                        "group_id": group_id,
                        "sheet": worksheet.title,
                        "rows": [],
                        "issues": [],
                        "fields": {},
                        "items": [],
                        "evidence": [],
                    },
                )
                row_data: dict[str, Any] = {}
                row_locations: dict[str, dict[str, object]] = {}
                for field, column in headers.items():
                    cell = row[column - 1] if column <= len(row) else None
                    value = cell.value if cell is not None else None
                    location = {
                        "sheet": worksheet.title,
                        "row": row[0].row,
                        "column": column,
                        "cell": f"{get_column_letter(column)}{row[0].row}",
                    }
                    row_locations[field] = location
                    if cell is not None and cell.data_type == "f":
                        group["issues"].append(
                            _cell_issue(
                                "FORMULA_NOT_ALLOWED",
                                field,
                                "业务字段不能使用公式，请替换为固定值。",
                                location,
                            )
                        )
                        continue
                    row_data[field] = value
                if not group_id.startswith("（缺少分组）"):
                    group["evidence"].append(
                        FieldEvidenceInput(
                            field_path=f"groups.{worksheet.title}.{group_id}.group_id",
                            source_type=FieldSource.FILE_CELL,
                            artifact_id=artifact_id,
                            locator=row_locations.get("group_id"),
                            raw_value=raw_group,
                            accepted_value=group_id,
                        )
                    )
                else:
                    group["issues"].append(
                        _cell_issue(
                            "MISSING_FIELD",
                            "group_id",
                            "请填写订单分组编号；同一订单的多条明细使用相同编号。",
                            row_locations["group_id"],
                        )
                    )
                group["rows"].append(row[0].row)
                for field in _GROUP_FIELDS:
                    raw = row_data.get(field)
                    try:
                        normalized: Any = (
                            _order_type(raw)
                            if field == "order_type"
                            else (
                                _date_text(raw)
                                if field == "delivery_date"
                                else _text(raw)
                            )
                        )
                    except ValueError as exc:
                        group["issues"].append(
                            _cell_issue(
                                "INVALID_CELL_VALUE",
                                field,
                                str(exc),
                                row_locations[field],
                            )
                        )
                        continue
                    if normalized is None:
                        if field == "order_type" and _text(raw) is not None:
                            group["issues"].append(
                                _cell_issue(
                                    "INVALID_ORDER_TYPE",
                                    field,
                                    "订单类型只能填写销售或采购。",
                                    row_locations[field],
                                )
                            )
                            continue
                        group["issues"].append(
                            _cell_issue(
                                "MISSING_FIELD",
                                field,
                                "该字段不能为空。",
                                row_locations[field],
                            )
                        )
                        continue
                    previous = group["fields"].get(field)
                    if previous is not None and previous[0] != normalized:
                        group["issues"].append(
                            _cell_issue(
                                "GROUP_FIELD_CONFLICT",
                                field,
                                "同一订单组的主档字段不一致，请统一后再导入。",
                                row_locations[field],
                            )
                        )
                        group["issues"].append(
                            _cell_issue(
                                "GROUP_FIELD_CONFLICT",
                                field,
                                "该组其他行使用了不同值。",
                                previous[1],
                            )
                        )
                    else:
                        group["fields"][field] = (normalized, row_locations[field])
                    group["evidence"].append(
                        FieldEvidenceInput(
                            field_path=field,
                            source_type=FieldSource.FILE_CELL,
                            artifact_id=artifact_id,
                            locator=row_locations[field],
                            raw_value=raw,
                            accepted_value=(
                                normalized.isoformat()
                                if isinstance(normalized, date)
                                else normalized
                            ),
                        )
                    )

                for field in ("raw_name", "specification", "unit"):
                    value = _text(row_data.get(field))
                    if value is None:
                        group["issues"].append(
                            _cell_issue(
                                "MISSING_FIELD",
                                field,
                                "该字段不能为空。",
                                row_locations[field],
                            )
                        )
                quantity = None
                price = None
                for field in ("quantity", "unit_price"):
                    try:
                        number = _decimal_text(row_data.get(field))
                    except ValueError as exc:
                        group["issues"].append(
                            _cell_issue(
                                "INVALID_NUMBER", field, str(exc), row_locations[field]
                            )
                        )
                        continue
                    if field == "quantity":
                        quantity = number
                    else:
                        price = number
                    if number is None or Decimal(number) <= 0:
                        code = (
                            "INVALID_QUANTITY"
                            if field == "quantity"
                            else "INVALID_UNIT_PRICE"
                        )
                        label = "数量" if field == "quantity" else "含税单价"
                        group["issues"].append(
                            _cell_issue(
                                code,
                                field,
                                f"{label}必须大于 0。",
                                row_locations[field],
                            )
                        )
                group["items"].append(
                    {
                        "material_id": None,
                        "raw_name": _text(row_data.get("raw_name")),
                        "specification": _text(row_data.get("specification")),
                        "quantity": quantity,
                        "unit": _text(row_data.get("unit")),
                        "unit_price": price,
                    }
                )
                for field in (
                    "raw_name",
                    "specification",
                    "quantity",
                    "unit",
                    "unit_price",
                ):
                    value = row_data.get(field)
                    if value is None or field not in row_data:
                        continue
                    group["evidence"].append(
                        FieldEvidenceInput(
                            field_path=f"items[{len(group['items']) - 1}].{field}",
                            source_type=FieldSource.FILE_CELL,
                            artifact_id=artifact_id,
                            locator=row_locations[field],
                            raw_value=value,
                            accepted_value=(
                                quantity
                                if field == "quantity"
                                else price
                                if field == "unit_price"
                                else _text(value)
                            ),
                        )
                    )

        for group in groups.values():
            contract_number = group["fields"].get("contract_number", (None, None))[0]
            contract_row = None
            if contract_number is not None:
                if contract_number not in contract_cache:
                    contract_cache[contract_number] = await _find_contract(
                        session, department_ids, contract_number
                    )
                contract_row = contract_cache[contract_number]
            if contract_row is None:
                location = group["fields"].get(
                    "contract_number",
                    (
                        None,
                        {
                            "sheet": group["sheet"],
                            "row": group["rows"][0],
                            "column": None,
                            "cell": None,
                        },
                    ),
                )[1]
                group["issues"].append(
                    _cell_issue(
                        "CONTRACT_NOT_FOUND",
                        "contract_number",
                        "当前部门找不到唯一有效合同，请检查合同编号。",
                        location,
                    )
                )
            else:
                contract, customer, project = contract_row
                if contract.status != "ACTIVE":
                    group["issues"].append(
                        _cell_issue(
                            "CONTRACT_INACTIVE",
                            "contract_number",
                            "合同当前未生效，不能用于导入。",
                            group["fields"]["contract_number"][1],
                        )
                    )
                delivery_date = group["fields"].get("delivery_date", (None,))[0]
                if delivery_date is not None and not (
                    contract.valid_from <= delivery_date <= contract.valid_to
                ):
                    group["issues"].append(
                        _cell_issue(
                            "CONTRACT_DATE_CONFLICT",
                            "delivery_date",
                            "交期不在合同有效日期范围内。",
                            group["fields"]["delivery_date"][1],
                        )
                    )
                if (
                    "order_type" in group["fields"]
                    and group["fields"]["order_type"][0] != contract.type
                ):
                    group["issues"].append(
                        _cell_issue(
                            "CONTRACT_TYPE_CONFLICT",
                            "order_type",
                            "订单类型与合同类型不一致。",
                            group["fields"]["order_type"][1],
                        )
                    )
                if (
                    _text(group["fields"].get("project", (None,))[0])
                    not in {
                        _text(project.name),
                        _text(project.code),
                    }
                    and "project" in group["fields"]
                ):
                    group["issues"].append(
                        _cell_issue(
                            "PROJECT_CONTRACT_CONFLICT",
                            "project",
                            "项目与合同关联项目不一致。",
                            group["fields"]["project"][1],
                        )
                    )
                group["contract_id"] = contract.id
                group["customer_id"] = customer.id
                group["project_id"] = project.id
            fields = group["fields"]
            candidate = CandidateOrder(
                order_type=fields.get("order_type", (None,))[0],
                contract_id=group.get("contract_id"),
                customer_id=group.get("customer_id"),
                project_id=group.get("project_id"),
                delivery_date=fields.get("delivery_date", (None,))[0],
                address=fields.get("address", (None,))[0],
                contact=fields.get("contact", (None,))[0],
                items=tuple(CandidateOrderItem(**item) for item in group["items"]),
                evidence=tuple(group["evidence"]),
            )
            try:
                _, group_amount = candidate_amounts(candidate)
            except ValueError:
                group_amount = None
                group["issues"].append(
                    _cell_issue(
                        "AMOUNT_OUT_OF_RANGE",
                        "items",
                        "订单金额超出系统可处理范围，请检查数量和单价。",
                        {
                            "sheet": group["sheet"],
                            "row": group["rows"][0],
                            "column": None,
                            "cell": None,
                        },
                    )
                )
            group["total_amount"] = (
                str(group_amount) if group_amount is not None else None
            )
            quota_exceeded = False
            contract_id = group.get("contract_id")
            if contract_id is not None and group_amount is not None:
                if contract_id not in balance_cache:
                    balance_cache[contract_id] = await session.get(
                        ContractBalance, contract_id
                    )
                balance = balance_cache[contract_id]
                if balance is None:
                    group["issues"].append(
                        _cell_issue(
                            "QUOTA_UNAVAILABLE",
                            "contract_number",
                            "合同额度数据暂不可用，当前仅显示候选草稿。",
                            group["fields"]["contract_number"][1],
                        )
                    )
                else:
                    already_grouped = imported_by_contract.get(contract_id, Decimal(0))
                    remaining = max(
                        Decimal(0),
                        contract.total_amount
                        - balance.occupied_amount
                        - already_grouped,
                    )
                    quota_exceeded = group_amount > remaining
                    if not group["issues"]:
                        imported_by_contract[contract_id] = (
                            already_grouped + group_amount
                        )
                    group["quota"] = {
                        "contract_amount": str(contract.total_amount),
                        "occupied_amount": str(balance.occupied_amount),
                        "remaining_before_group": str(remaining),
                        "requested_amount": str(group_amount),
                        "over_quota": quota_exceeded,
                    }
            group["candidate"] = candidate.model_dump(mode="json")
            group["status"] = (
                "NEEDS_CORRECTION"
                if group["issues"]
                else "OVER_QUOTA_REVIEW"
                if quota_exceeded
                else "READY_FOR_REVIEW"
            )
            del group["fields"]
            del group["items"]
        if row_count == 0 and not issues:
            issues.append(
                {
                    "code": "EMPTY_IMPORT",
                    "field": "items",
                    "message": "Excel 模板中没有可导入的订单明细。",
                    "locator": {
                        "sheet": None,
                        "row": None,
                        "column": None,
                        "cell": None,
                    },
                }
            )
        return {
            "groups": list(groups.values()),
            "issues": issues,
            "summary": {
                "sheet_count": len(workbook.sheetnames),
                "detail_rows": row_count,
                "group_count": len(groups),
                "ready_groups": sum(
                    group["status"] == "READY_FOR_REVIEW" for group in groups.values()
                ),
                "over_quota_groups": sum(
                    group["status"] == "OVER_QUOTA_REVIEW" for group in groups.values()
                ),
                "needs_correction_groups": sum(
                    group["status"] == "NEEDS_CORRECTION" for group in groups.values()
                ),
            },
            "source_refs": [str(artifact_id)],
        }
    finally:
        workbook.close()
        source.close()
