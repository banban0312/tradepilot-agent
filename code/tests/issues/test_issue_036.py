"""M6-01 固定评测数据集契约与夹具完整性检查。"""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "evals"
DATASET = ROOT / "datasets" / "v1"
FIXTURES = ROOT / "fixtures" / "v1"
REQUIRED_CATEGORIES = {
    "normal",
    "abnormal",
    "ambiguity",
    "unauthorized",
    "quota",
    "duplicate_submission",
}
MODEL_ALLOWED_TOOLS = {
    "query_customer",
    "match_material",
    "get_contract",
    "check_contract_balance",
    "save_order_draft",
}


def _load_cases() -> list[dict]:
    return [
        json.loads(line)
        for line in (DATASET / "cases.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def test_dataset_has_six_categories_and_dedicated_recovery_concurrency_cases() -> None:
    manifest = json.loads((DATASET / "manifest.json").read_text(encoding="utf-8"))
    cases = _load_cases()
    counts = Counter(case["category"] for case in cases)

    assert manifest["dataset_version"] == "1.0.0"
    assert manifest["timezone"] == "Asia/Hong_Kong"
    assert manifest["fixed_clock"] == "2026-10-16T09:00:00+08:00"
    assert manifest["case_count"] == len(cases) >= 36
    assert set(manifest["required_categories"]) == REQUIRED_CATEGORIES
    assert all(counts[category] >= 6 for category in REQUIRED_CATEGORIES)
    assert sum("recovery" in case["tags"] for case in cases) >= 2
    assert sum("concurrency" in case["tags"] for case in cases) >= 2
    assert sum("unknown_commit_result" in case["tags"] for case in cases) >= 1
    assert manifest["category_counts"] == {
        category: counts[category] for category in REQUIRED_CATEGORIES
    }


def test_each_case_has_machine_assertions_and_references_known_fixtures() -> None:
    cases = _load_cases()
    fixture_index = json.loads((FIXTURES / "index.json").read_text(encoding="utf-8"))
    actor_ids = set(fixture_index["actors"])
    contract_ids = set(fixture_index["contracts"])

    assert len({case["id"] for case in cases}) == len(cases)
    for case in cases:
        assert re.fullmatch(r"[a-z0-9][a-z0-9-]+", case["id"])
        assert case["dataset_version"] == "1.0.0"
        assert case["master_data_version"] == "1.0.0"
        assert case["input_fixture_version"] == "1.0.0"
        assert case["fixed_time"] == "2026-10-16T09:00:00+08:00"
        assert set(case["execution_modes"]) == {"offline", "online"}
        assert datetime.fromisoformat(case["fixed_time"]).utcoffset() is not None
        assert case["actor_fixture"] in actor_ids
        assert case["contract_fixture"] in contract_ids
        assert case["preconditions"]["task_state"]
        assert case["input"].get("value") or case["input"].get("fixture_id")
        assert case["actions"]
        expected = case["expected"]
        assert expected["task_state"]
        assert "error_code" in expected
        assert "required_tools" in expected and "allowed_tools" in expected
        assert set(expected["allowed_tools"]) == MODEL_ALLOWED_TOOLS
        assert set(expected["required_tools"]).issubset(expected["allowed_tools"])
        assert expected["assertions"]
        assert all(
            set(assertion) == {"source", "operator", "expected"}
            and assertion["operator"] in {"eq", "contains_all"}
            and isinstance(assertion["source"], str)
            for assertion in expected["assertions"]
        )
        assert expected["rules"]
        assert set(expected["facts"]) >= {
            "orders",
            "approvals",
            "quota_occupations",
            "occupied_amount_delta",
        }
        if case["input"].get("fixture_id"):
            assert case["input"]["fixture_id"] in fixture_index["files"]


def test_file_fixtures_match_manifest_hashes_and_contain_no_credentials() -> None:
    fixture_index = json.loads((FIXTURES / "index.json").read_text(encoding="utf-8"))
    manifest = json.loads((DATASET / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["files"] == fixture_index["files"]
    for fixture in fixture_index["files"].values():
        path = FIXTURES / fixture["path"]
        assert path.is_file()
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        assert digest == fixture["sha256"]

    workbook = load_workbook(
        FIXTURES / fixture_index["files"]["mixed_batch"]["path"],
        read_only=True,
        data_only=True,
    )
    sheet = workbook["订单明细"]
    assert tuple(cell.value for cell in next(sheet.iter_rows(max_row=1))) == (
        "订单分组",
        "订单类型",
        "合同编号",
        "项目",
        "地址",
        "联系人",
        "交期",
        "物料名称",
        "规格",
        "数量",
        "单位",
        "含税单价",
    )
    assert sheet.max_row == 4
    workbook.close()
    assert (
        (FIXTURES / fixture_index["files"]["damaged_pdf"]["path"])
        .read_bytes()
        .startswith(b"%PDF-1.4")
    )

    serialized = (
        "\n".join(
            path.read_text(encoding="utf-8")
            for path in [
                DATASET / "manifest.json",
                FIXTURES / "master_data.json",
                FIXTURES / "index.json",
            ]
        )
        + "\n"
        + "\n".join((DATASET / "cases.jsonl").read_text(encoding="utf-8").splitlines())
    )
    assert not re.search(
        r"sk-proj-[A-Za-z0-9_-]{20,}|sk-[A-Za-z0-9_-]{24,}", serialized
    )
    assert not re.search(
        r"(?i)(password|passwd|api[_-]?key)\s*[:=]\s*['\"]?[^\s,}\"]{8,}", serialized
    )
