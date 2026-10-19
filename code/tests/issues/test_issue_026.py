"""M4-05 Excel 模板、订单分组和单元格校验验收。"""

from __future__ import annotations

import io
import os
import time
from contextlib import asynccontextmanager
from uuid import UUID

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.database import async_session_factory
from app.models import AuditEvent, Order

API_URL = os.getenv("TRADEPILOT_TEST_API_URL", "http://api:8000")
ORIGIN = os.getenv("APP_ORIGIN", "http://localhost:5173")
HEADERS = [
    "订单分组",
    "订单类型",
    "合同编号",
    "项目",
    "地址",
    "联系人",
    "交期",
    "物料名称",
    "规格",
    "数量",
    "单位",
    "含税单价",
]


def _workbook(rows: list[list[object]], *, formulas: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单明细"
    sheet.append(HEADERS)
    for values in rows:
        sheet.append(values)
    if formulas:
        sheet["L2"] = "=100+200"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _row(group: str, name: str, quantity: object, price: object = 100) -> list[object]:
    return [
        group,
        "销售",
        "TP-DEMO-SALES-N1",
        "华北示例工程",
        "北京市朝阳区示例路 1 号",
        "张先生",
        "2026-11-01",
        name,
        "Q235 8mm",
        quantity,
        "吨",
        price,
    ]


@asynccontextmanager
async def _client(login: str = "demo-operator-north"):
    assert os.getenv("TRADEPILOT_TEST_MODE") == "1", "请通过隔离 Compose 验收脚本运行"
    password = os.getenv("DEMO_PASSWORD", "")
    assert len(password) >= 12
    client = httpx.AsyncClient(base_url=API_URL, timeout=30)
    response = await client.post(
        "/api/v1/auth/login",
        headers={"Origin": ORIGIN},
        json={"login": login, "password": password},
    )
    assert response.status_code == 200, response.text
    client.headers.update(
        {"Origin": ORIGIN, "X-CSRF-Token": response.json()["data"]["csrf_token"]}
    )
    try:
        yield client
    finally:
        await client.aclose()


async def _upload(client: httpx.AsyncClient, content: bytes) -> UUID:
    response = await client.post(
        "/api/v1/artifacts",
        data={"parse_intent": "order_xlsx"},
        files={"file": ("订单导入.xlsx", content, "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    return UUID(response.json()["data"]["artifact_id"])


@pytest.mark.asyncio
async def test_template_download_groups_rows_and_locates_conflicts_and_quota() -> None:
    async with _client() as client:
        template = await client.get("/api/v1/imports/template")
        assert template.status_code == 200
        workbook = load_workbook(io.BytesIO(template.content), read_only=True)
        assert tuple(
            cell.value for cell in next(workbook["订单明细"].iter_rows())
        ) == tuple(HEADERS)
        assert "填报说明" in workbook.sheetnames
        workbook.close()
        downloaded_template = await _upload(client, template.content)
        template_preview = await client.post(
            "/api/v1/imports/preview",
            json={"artifact_id": str(downloaded_template)},
        )
        assert template_preview.status_code == 200, template_preview.text
        assert template_preview.json()["data"]["summary"]["detail_rows"] == 0
        assert template_preview.json()["data"]["issues"][0]["code"] == "EMPTY_IMPORT"

        rows = [
            _row("正常组", "钢板 A", 2),
            _row("正常组", "钢板 B", 3),
            _row("超额组", "大量钢板", 2000),
            _row("冲突组", "冲突钢板", 1),
            _row("冲突组", "另一种钢板", 1),
        ]
        rows[-1][4] = "其他地址"
        artifact_id = await _upload(client, _workbook(rows))
        async with async_session_factory() as session:
            before_orders = await session.scalar(
                select(func.count()).select_from(Order)
            )

        preview = await client.post(
            "/api/v1/imports/preview", json={"artifact_id": str(artifact_id)}
        )
        assert preview.status_code == 200, preview.text
        data = preview.json()["data"]
        assert data["summary"] == {
            "sheet_count": 1,
            "detail_rows": 5,
            "group_count": 3,
            "ready_groups": 1,
            "over_quota_groups": 1,
            "needs_correction_groups": 1,
        }
        groups = {group["group_id"]: group for group in data["groups"]}
        normal = groups["正常组"]
        assert normal["status"] == "READY_FOR_REVIEW"
        assert normal["total_amount"] == "500.00"
        assert len(normal["candidate"]["items"]) == 2
        assert normal["candidate"]["items"][0]["material_id"] is None
        evidence = [
            row
            for row in normal["candidate"]["evidence"]
            if row["field_path"] == "items[1].raw_name"
        ]
        assert evidence[0]["artifact_id"] == str(artifact_id)
        assert evidence[0]["locator"] == {
            "sheet": "订单明细",
            "row": 3,
            "column": 8,
            "cell": "H3",
        }
        assert groups["超额组"]["status"] == "OVER_QUOTA_REVIEW"
        assert groups["超额组"]["quota"]["over_quota"] is True
        conflict = groups["冲突组"]
        assert conflict["status"] == "NEEDS_CORRECTION"
        assert any(
            issue["code"] == "GROUP_FIELD_CONFLICT" and issue["locator"]["cell"] == "E6"
            for issue in conflict["issues"]
        )
        async with async_session_factory() as session:
            after_orders = await session.scalar(select(func.count()).select_from(Order))
            audit = await session.scalar(
                select(AuditEvent).where(
                    AuditEvent.action == "import.previewed",
                    AuditEvent.object_ref["artifact_id"].astext == str(artifact_id),
                )
            )
        assert after_orders == before_orders
        assert audit is not None and audit.input_hash

        async with _client("demo-operator-south") as other_department:
            denied = await other_department.post(
                "/api/v1/imports/preview", json={"artifact_id": str(artifact_id)}
            )
            assert denied.status_code == 404


@pytest.mark.asyncio
async def test_import_rejects_formulas_bad_numbers_missing_columns_and_excess_rows() -> (
    None
):
    async with _client() as client:
        formula_artifact = await _upload(
            client,
            _workbook([_row("公式组", "钢板", 1)], formulas=True),
        )
        formula = await client.post(
            "/api/v1/imports/preview", json={"artifact_id": str(formula_artifact)}
        )
        assert formula.status_code == 200, formula.text
        formula_group = formula.json()["data"]["groups"][0]
        assert formula_group["status"] == "NEEDS_CORRECTION"
        assert any(
            issue["code"] == "FORMULA_NOT_ALLOWED" and issue["locator"]["cell"] == "L2"
            for issue in formula_group["issues"]
        )

        invalid_artifact = await _upload(
            client,
            _workbook([_row("数字组", "钢板", "非数字", "1.1234567")]),
        )
        invalid = await client.post(
            "/api/v1/imports/preview", json={"artifact_id": str(invalid_artifact)}
        )
        assert invalid.status_code == 200, invalid.text
        number_issues = invalid.json()["data"]["groups"][0]["issues"]
        assert {
            issue["field"]
            for issue in number_issues
            if issue["code"] == "INVALID_NUMBER"
        } == {"quantity", "unit_price"}

        missing_columns = Workbook()
        missing_columns.active.append(["订单分组", "物料名称"])
        output = io.BytesIO()
        missing_columns.save(output)
        missing_columns.close()
        missing_artifact = await _upload(client, output.getvalue())
        missing = await client.post(
            "/api/v1/imports/preview", json={"artifact_id": str(missing_artifact)}
        )
        assert missing.status_code == 200, missing.text
        assert all(
            issue["code"] == "MISSING_COLUMN"
            for issue in missing.json()["data"]["issues"]
        )

        empty_artifact = await _upload(client, _workbook([]))
        empty = await client.post(
            "/api/v1/imports/preview", json={"artifact_id": str(empty_artifact)}
        )
        assert empty.status_code == 200, empty.text
        assert empty.json()["data"]["issues"][0]["code"] == "EMPTY_IMPORT"

        date_row = _row("合同日期组", "钢板", 1)
        date_row[6] = "2031-01-01"
        date_artifact = await _upload(client, _workbook([date_row]))
        date_preview = await client.post(
            "/api/v1/imports/preview", json={"artifact_id": str(date_artifact)}
        )
        assert date_preview.status_code == 200, date_preview.text
        assert any(
            issue["code"] == "CONTRACT_DATE_CONFLICT"
            for issue in date_preview.json()["data"]["groups"][0]["issues"]
        )

        hundred_rows = await _upload(
            client,
            _workbook([_row("百行组", f"钢板-{index}", 1) for index in range(100)]),
        )
        started = time.perf_counter()
        hundred_row_preview = await client.post(
            "/api/v1/imports/preview", json={"artifact_id": str(hundred_rows)}
        )
        elapsed = time.perf_counter() - started
        assert hundred_row_preview.status_code == 200, hundred_row_preview.text
        assert hundred_row_preview.json()["data"]["summary"]["detail_rows"] == 100
        assert elapsed < 10, f"100 行 XLSX 预览耗时 {elapsed:.2f} 秒，超过 10 秒目标"

        excess = await client.post(
            "/api/v1/artifacts",
            data={"parse_intent": "order_xlsx"},
            files={
                "file": (
                    "超行数.xlsx",
                    _workbook([_row(f"组{i}", "钢板", 1) for i in range(501)]),
                    "application/octet-stream",
                )
            },
        )
        assert excess.status_code == 413, excess.text
        assert excess.json()["error"]["code"] == "FILE_TOO_LARGE"
